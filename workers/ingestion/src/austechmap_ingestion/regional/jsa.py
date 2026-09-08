"""Versioned Jobs and Skills Australia NERO and IVI imports.

The two official products do not share one directly comparable occupation/
geography grain. NERO is ANZSCO-4 by SA4. The downloadable IVI SA4 workbook is
ANZSCO-2 and represents capital-city labour markets as GCCSAs. This module
preserves those distinctions and never manufactures SA4 rows for a GCCSA.
"""

from __future__ import annotations

import calendar
import csv
import hashlib
import logging
import time
import uuid
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import TextIOWrapper
from pathlib import Path, PurePosixPath
from typing import Literal, cast

import psycopg
from openpyxl import load_workbook

from austechmap_ingestion.jobs import ClaimedRun, JobRepository, SnapshotRecord
from austechmap_ingestion.storage import SnapshotStore

Dataset = Literal["nero", "ivi"]

NERO_SOURCE_KEY = "jsa-nero"
NERO_SOURCE_URL = "https://www.jobsandskills.gov.au/data/nero"
NERO_ATTRIBUTION = (
    "Nowcast of Employment by Region and Occupation, Jobs and Skills Australia, "
    "Commonwealth of Australia. Used under Creative Commons BY 4.0 licence."
)
IVI_SOURCE_KEY = "jsa-internet-vacancy-index"
IVI_SOURCE_URL = "https://www.jobsandskills.gov.au/data/internet-vacancy-index"
IVI_ATTRIBUTION = "© Commonwealth of Australia"

DEFAULT_MAPPING_PATH = Path(__file__).parent / "fixtures" / "anzsco4_role_family_v1.csv"
MAX_NERO_CSV_BYTES = 450 * 1024 * 1024
MAX_NERO_COMPRESSION_RATIO = 20
IVI_TECH_ANZSCO2_CODE = "26"
IVI_TECH_ANZSCO2_TITLE = "ICT Professionals"
CC_BY_4_URL = "https://creativecommons.org/licenses/by/4.0/"


class JsaImportError(RuntimeError):
    """Raised when an input release conflicts with its declared contract."""


@dataclass(frozen=True)
class OccupationRoleMapping:
    anzsco4_code: str
    anzsco4_title: str
    role_family_key: str
    rationale: str


@dataclass(frozen=True)
class RegionalLaborRecord:
    dataset: Dataset
    region_code: str
    region_name: str
    role_family_key: str | None
    metric_key: str
    period_start: date
    period_end: date
    value: Decimal
    unit: str
    direction: int | None = None


@dataclass(frozen=True)
class RegionalLaborLoadResult:
    records_seen: int
    records_inserted: int
    records_unchanged: int


@dataclass(frozen=True)
class JsaImportResult:
    run_id: uuid.UUID
    dataset: Dataset
    records_seen: int
    records_inserted: int
    records_unchanged: int
    created: bool


def load_occupation_role_mapping(
    path: Path = DEFAULT_MAPPING_PATH,
) -> dict[str, OccupationRoleMapping]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"anzsco4_code", "anzsco4_title", "role_family_key", "rationale"}
        if reader.fieldnames is None or set(reader.fieldnames) != required:
            raise JsaImportError(f"occupation mapping must have columns {sorted(required)}")
        mappings: dict[str, OccupationRoleMapping] = {}
        for row in reader:
            mapping = OccupationRoleMapping(
                anzsco4_code=row["anzsco4_code"].strip(),
                anzsco4_title=row["anzsco4_title"].strip(),
                role_family_key=row["role_family_key"].strip(),
                rationale=row["rationale"].strip(),
            )
            if (
                len(mapping.anzsco4_code) != 4
                or not mapping.anzsco4_code.isdigit()
                or not mapping.anzsco4_title
                or not mapping.role_family_key
                or not mapping.rationale
            ):
                raise JsaImportError(f"invalid occupation mapping row for {mapping.anzsco4_code!r}")
            if mapping.anzsco4_code in mappings:
                raise JsaImportError(f"duplicate ANZSCO-4 mapping {mapping.anzsco4_code}")
            mappings[mapping.anzsco4_code] = mapping
    if not mappings:
        raise JsaImportError("occupation mapping is empty")
    return mappings


def _month_bounds(value: date) -> tuple[date, date]:
    start = value.replace(day=1)
    return start, start.replace(day=calendar.monthrange(start.year, start.month)[1])


def _value(raw: object, *, context: str) -> Decimal:
    try:
        value = Decimal(str(raw).strip())
    except (InvalidOperation, ValueError) as error:
        raise JsaImportError(f"{context} has a non-numeric value {raw!r}") from error
    if not value.is_finite() or value < 0:
        raise JsaImportError(f"{context} must be a finite, non-negative value")
    return value


def _record_key(record: RegionalLaborRecord) -> tuple[object, ...]:
    return (
        record.region_code,
        record.role_family_key,
        record.dataset,
        record.metric_key,
        record.period_start,
        record.period_end,
    )


def _require_unique(records: list[RegionalLaborRecord]) -> list[RegionalLaborRecord]:
    keys: set[tuple[object, ...]] = set()
    for record in records:
        key = _record_key(record)
        if key in keys:
            raise JsaImportError(f"duplicate regional labour record {key}")
        keys.add(key)
    if not records:
        raise JsaImportError("release produced no supported regional labour records")
    return records


def parse_nero_zip(
    path: Path,
    *,
    mappings: dict[str, OccupationRoleMapping] | None = None,
) -> list[RegionalLaborRecord]:
    role_mappings = mappings if mappings is not None else load_occupation_role_mapping()
    required_columns = {
        "state_name",
        "sa4_code",
        "sa4_name",
        "anzsco4_code",
        "anzsco4_name",
        "date",
        "nsc_emp",
    }
    records: list[RegionalLaborRecord] = []
    try:
        with zipfile.ZipFile(path) as archive:
            csv_entries = [entry for entry in archive.infolist() if entry.filename.endswith(".csv")]
            if len(csv_entries) != 1:
                raise JsaImportError("NERO archive must contain exactly one CSV")
            entry = csv_entries[0]
            parts = PurePosixPath(entry.filename).parts
            if entry.filename.startswith("/") or ".." in parts:
                raise JsaImportError("NERO archive contains an unsafe CSV path")
            if entry.file_size > MAX_NERO_CSV_BYTES:
                raise JsaImportError("NERO CSV exceeds the configured uncompressed-size limit")
            compression_ratio = (
                entry.file_size / entry.compress_size if entry.compress_size else float("inf")
            )
            if compression_ratio > MAX_NERO_COMPRESSION_RATIO:
                raise JsaImportError("NERO CSV exceeds the configured compression-ratio limit")

            with (
                archive.open(entry) as binary_handle,
                TextIOWrapper(binary_handle, encoding="utf-8-sig", newline="") as text_handle,
            ):
                reader = csv.DictReader(text_handle)
                if reader.fieldnames is None or not required_columns <= set(reader.fieldnames):
                    raise JsaImportError(
                        f"NERO CSV is missing columns {sorted(required_columns)}"
                    )
                for row_number, row in enumerate(reader, start=2):
                    occupation_code = row["anzsco4_code"].strip()
                    mapping = role_mappings.get(occupation_code)
                    if mapping is None:
                        continue
                    occupation_title = row["anzsco4_name"].strip()
                    if occupation_title != mapping.anzsco4_title:
                        raise JsaImportError(
                            f"NERO row {row_number} title for {occupation_code} changed from "
                            f"{mapping.anzsco4_title!r} to {occupation_title!r}"
                        )
                    try:
                        source_date = date.fromisoformat(row["date"].strip())
                    except ValueError as error:
                        raise JsaImportError(
                            f"NERO row {row_number} has an invalid date"
                        ) from error
                    period_start, period_end = _month_bounds(source_date)
                    region_code = row["sa4_code"].strip()
                    region_name = row["sa4_name"].strip()
                    if not region_code.isdigit() or not region_name:
                        raise JsaImportError(f"NERO row {row_number} has invalid SA4 metadata")
                    records.append(
                        RegionalLaborRecord(
                            dataset="nero",
                            region_code=region_code,
                            region_name=region_name,
                            role_family_key=mapping.role_family_key,
                            metric_key=f"employment_anzsco4_{occupation_code}",
                            period_start=period_start,
                            period_end=period_end,
                            value=_value(row["nsc_emp"], context=f"NERO row {row_number}"),
                            unit="employed_persons_nowcast",
                        )
                    )
    except (OSError, zipfile.BadZipFile) as error:
        raise JsaImportError(f"could not read NERO archive {path}") from error
    return _require_unique(records)


def parse_ivi_workbook(path: Path) -> list[RegionalLaborRecord]:
    expected_prefix = (
        "Level",
        "State",
        "region_name",
        "region_code",
        "region_level",
        "ANZSCO_CODE",
        "ANZSCO_TITLE",
    )
    records: list[RegionalLaborRecord] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as error:
        raise JsaImportError(f"could not read IVI workbook {path}") from error
    try:
        if "Averaged" not in workbook.sheetnames:
            raise JsaImportError("IVI workbook is missing the Averaged sheet")
        worksheet = workbook["Averaged"]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header[:7]) != expected_prefix:
            raise JsaImportError("IVI Averaged sheet metadata columns changed")
        date_columns: list[tuple[int, date, date]] = []
        for index, raw_date in enumerate(header[7:], start=7):
            if isinstance(raw_date, datetime):
                month = raw_date.date()
            elif isinstance(raw_date, date):
                month = raw_date
            else:
                raise JsaImportError(f"IVI column {index + 1} is not a monthly date")
            period_start, period_end = _month_bounds(month)
            date_columns.append((index, period_start, period_end))

        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None for value in row):
                continue
            if len(row) < len(expected_prefix):
                raise JsaImportError(f"IVI row {row_number} is missing metadata columns")
            if str(row[4]).strip().upper() != "SA4":
                continue
            if str(row[5]).strip() != IVI_TECH_ANZSCO2_CODE:
                continue
            if str(row[6]).strip() != IVI_TECH_ANZSCO2_TITLE:
                raise JsaImportError(
                    f"IVI row {row_number} title for code 26 changed to {row[6]!r}"
                )
            region_code = str(row[3]).strip()
            region_name = str(row[2]).strip()
            if not region_code.isdigit() or not region_name:
                raise JsaImportError(f"IVI row {row_number} has invalid SA4 metadata")
            for index, period_start, period_end in date_columns:
                raw_value = row[index] if index < len(row) else None
                if raw_value is None or str(raw_value).strip() == "":
                    continue
                records.append(
                    RegionalLaborRecord(
                        dataset="ivi",
                        region_code=region_code,
                        region_name=region_name,
                        role_family_key=None,
                        metric_key="vacancies_anzsco2_26_3mma",
                        period_start=period_start,
                        period_end=period_end,
                        value=_value(raw_value, context=f"IVI row {row_number}"),
                        unit="online_job_advertisements_3_month_moving_average",
                    )
                )
    finally:
        workbook.close()
    return _require_unique(records)


def load_regional_labor_observations(
    database_url: str,
    *,
    records: list[RegionalLaborRecord],
    source_version: str,
    source_id: uuid.UUID,
    import_run_id: uuid.UUID,
    observed_at: datetime,
) -> RegionalLaborLoadResult:
    _require_unique(records)
    if not source_version.strip():
        raise ValueError("source_version cannot be empty")
    if observed_at.tzinfo is None or observed_at.utcoffset() is None:
        raise ValueError("observed_at must be timezone-aware")

    with psycopg.connect(database_url) as connection:
        connection.execute(
            """
            CREATE TEMP TABLE staged_regional_labor (
              dataset TEXT NOT NULL,
              region_code TEXT NOT NULL,
              role_family_key TEXT,
              metric_key TEXT NOT NULL,
              period_start DATE NOT NULL,
              period_end DATE NOT NULL,
              value NUMERIC NOT NULL,
              unit TEXT NOT NULL,
              direction SMALLINT
            ) ON COMMIT DROP
            """
        )
        with connection.cursor().copy(
            """
            COPY staged_regional_labor (
              dataset, region_code, role_family_key, metric_key, period_start,
              period_end, value, unit, direction
            ) FROM STDIN
            """
        ) as copy:
            for record in records:
                copy.write_row(
                    (
                        record.dataset,
                        record.region_code,
                        record.role_family_key,
                        record.metric_key,
                        record.period_start,
                        record.period_end,
                        record.value,
                        record.unit,
                        record.direction,
                    )
                )

        missing_regions = [
            cast(str, row[0])
            for row in connection.execute(
                """
                WITH active_sa4 AS (
                  SELECT r.code
                  FROM regions r
                  JOIN geography_releases gr ON gr.id = r.release_id
                  WHERE r.region_type = 'sa4'
                    AND gr.dataset = 'asgs_sa4'
                    AND gr.is_active
                )
                SELECT DISTINCT staged.region_code
                FROM staged_regional_labor staged
                LEFT JOIN active_sa4 region ON region.code = staged.region_code
                WHERE region.code IS NULL
                ORDER BY staged.region_code
                """
            ).fetchall()
        ]
        if missing_regions:
            raise JsaImportError(
                "release references SA4 codes absent from the active ASGS release: "
                f"{missing_regions}"
            )

        missing_roles = [
            cast(str, row[0])
            for row in connection.execute(
                """
                SELECT DISTINCT staged.role_family_key
                FROM staged_regional_labor staged
                LEFT JOIN role_families role ON role.key = staged.role_family_key
                WHERE staged.role_family_key IS NOT NULL AND role.id IS NULL
                ORDER BY staged.role_family_key
                """
            ).fetchall()
        ]
        if missing_roles:
            raise JsaImportError(
                f"occupation mapping references unknown role families: {missing_roles}"
            )

        conflict = connection.execute(
            """
            WITH active_sa4 AS (
              SELECT r.id, r.code
              FROM regions r
              JOIN geography_releases gr ON gr.id = r.release_id
              WHERE r.region_type = 'sa4'
                AND gr.dataset = 'asgs_sa4'
                AND gr.is_active
            )
            SELECT staged.region_code, staged.metric_key, staged.period_start
            FROM staged_regional_labor staged
            JOIN active_sa4 region ON region.code = staged.region_code
            LEFT JOIN role_families role ON role.key = staged.role_family_key
            JOIN regional_labor_observations existing
              ON existing.region_id = region.id
             AND existing.role_family_id IS NOT DISTINCT FROM role.id
             AND existing.dataset::text = staged.dataset
             AND existing.metric_key = staged.metric_key
             AND existing.period_start = staged.period_start
             AND existing.period_end = staged.period_end
             AND existing.source_version = %(source_version)s
            WHERE existing.value IS DISTINCT FROM staged.value
               OR existing.unit IS DISTINCT FROM staged.unit
               OR existing.direction IS DISTINCT FROM staged.direction
               OR existing.source_id IS DISTINCT FROM %(source_id)s
            LIMIT 1
            """,
            {"source_version": source_version, "source_id": source_id},
        ).fetchone()
        if conflict is not None:
            raise JsaImportError(
                "source version conflicts with an immutable observation for "
                f"SA4 {conflict[0]}, {conflict[1]}, {conflict[2]}"
            )

        inserted = connection.execute(
            """
            WITH active_sa4 AS (
              SELECT r.id, r.code
              FROM regions r
              JOIN geography_releases gr ON gr.id = r.release_id
              WHERE r.region_type = 'sa4'
                AND gr.dataset = 'asgs_sa4'
                AND gr.is_active
            )
            INSERT INTO regional_labor_observations (
              region_id, role_family_id, dataset, metric_key, period_start,
              period_end, value, unit, direction, source_version, source_id,
              import_run_id, observed_at
            )
            SELECT
              region.id, role.id, staged.dataset::regional_labor_dataset,
              staged.metric_key, staged.period_start, staged.period_end,
              staged.value, staged.unit, staged.direction, %(source_version)s,
              %(source_id)s, %(import_run_id)s, %(observed_at)s
            FROM staged_regional_labor staged
            JOIN active_sa4 region ON region.code = staged.region_code
            LEFT JOIN role_families role ON role.key = staged.role_family_key
            ON CONFLICT (
              region_id, role_family_id, dataset, metric_key, period_start,
              period_end, source_version
            ) DO NOTHING
            RETURNING id
            """,
            {
                "source_version": source_version,
                "source_id": source_id,
                "import_run_id": import_run_id,
                "observed_at": observed_at,
            },
        ).fetchall()

    return RegionalLaborLoadResult(
        records_seen=len(records),
        records_inserted=len(inserted),
        records_unchanged=len(records) - len(inserted),
    )


def _finalize_failed_import(
    repository: JobRepository,
    claim: ClaimedRun,
    error: Exception,
    *,
    now: datetime,
) -> None:
    try:
        repository.fail(
            claim,
            retryable=True,
            error_code=type(error).__name__,
            error_message=str(error),
            jitter_fraction=0,
            now=now,
        )
    except Exception:
        logging.exception("could not record failed JSA import %s", claim.run_id)


def _configure_jsa_source(
    database_url: str,
    *,
    source_id: uuid.UUID,
    source_url: str,
    attribution_text: str,
) -> None:
    with psycopg.connect(database_url) as connection:
        updated = connection.execute(
            """
            UPDATE data_sources
            SET base_url = %s,
                licence_name = 'Creative Commons Attribution 4.0 International',
                licence_url = %s,
                attribution_text = %s,
                retrieval_policy = '{"mode":"manual_official_download"}'::jsonb
            WHERE id = %s
            RETURNING id
            """,
            (
                source_url,
                CC_BY_4_URL,
                attribution_text,
                source_id,
            ),
        ).fetchone()
    if updated is None:
        raise JsaImportError(f"JSA source {source_id} disappeared during configuration")


def _run_jsa_import(
    repository: JobRepository,
    store: SnapshotStore,
    database_url: str,
    *,
    dataset: Dataset,
    source_key: str,
    source_name: str,
    source_url: str,
    source_attribution: str,
    source_version: str,
    file_path: Path,
    parser: Callable[[Path], list[RegionalLaborRecord]],
    content_type: str,
    worker_id: str,
    now: datetime | None,
) -> JsaImportResult:
    import_time = datetime.now(UTC) if now is None else now
    if import_time.tzinfo is None or import_time.utcoffset() is None:
        raise ValueError("now must be timezone-aware")
    source_version = source_version.strip()
    if not source_version:
        raise ValueError("source_version cannot be empty")
    content = file_path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    source_id = repository.ensure_source(
        source_key=source_key,
        name=source_name,
        kind="government_open_data",
    )
    _configure_jsa_source(
        database_url,
        source_id=source_id,
        source_url=source_url,
        attribution_text=source_attribution,
    )
    enqueued = repository.enqueue(
        run_type=f"jsa_{dataset}_import",
        idempotency_key=f"{dataset}:{source_version}:{digest}",
        source_id=source_id,
        payload={
            "dataset": dataset,
            "source_url": source_url,
            "source_version": source_version,
            "sha256": digest,
        },
        scheduled_for=import_time,
    )
    if not enqueued.created:
        return JsaImportResult(enqueued.run_id, dataset, 0, 0, 0, False)

    claim = repository.claim_run(enqueued.run_id, worker_id=worker_id, now=import_time)
    if claim is None:
        raise JsaImportError(f"new {dataset} import could not be claimed")
    started_at = time.monotonic()
    try:
        # Preserve the exact source before interpreting it, matching the
        # acquisition -> raw snapshot -> parsing pipeline contract.
        stored = store.put(source_key=source_key, content=content, content_type=content_type)
        records = parser(file_path)
        load_result = load_regional_labor_observations(
            database_url,
            records=records,
            source_version=source_version,
            source_id=source_id,
            import_run_id=claim.run_id,
            observed_at=import_time,
        )
        repository.complete_with_snapshot(
            claim,
            SnapshotRecord(
                source_id=source_id,
                object_key=stored.object_key,
                sha256=stored.sha256,
                content_type=content_type,
                byte_size=stored.byte_size,
                retrieved_at=import_time,
                response_metadata={
                    "importer": f"jsa-{dataset}",
                    "version": 1,
                    "source_version": source_version,
                    "source_url": source_url,
                    "records_seen": load_result.records_seen,
                    "records_inserted": load_result.records_inserted,
                },
            ),
            metrics={
                "records_seen": load_result.records_seen,
                "records_inserted": load_result.records_inserted,
                "records_unchanged": load_result.records_unchanged,
                "duration_ms": max(0, round((time.monotonic() - started_at) * 1000)),
            },
            now=import_time,
        )
    except Exception as error:
        _finalize_failed_import(repository, claim, error, now=import_time)
        raise

    return JsaImportResult(
        run_id=claim.run_id,
        dataset=dataset,
        records_seen=load_result.records_seen,
        records_inserted=load_result.records_inserted,
        records_unchanged=load_result.records_unchanged,
        created=True,
    )


def run_nero_import(
    repository: JobRepository,
    store: SnapshotStore,
    database_url: str,
    *,
    source_version: str,
    file_path: Path,
    worker_id: str = "jsa-nero-importer",
    now: datetime | None = None,
) -> JsaImportResult:
    return _run_jsa_import(
        repository,
        store,
        database_url,
        dataset="nero",
        source_key=NERO_SOURCE_KEY,
        source_name="Jobs and Skills Australia NERO",
        source_url=NERO_SOURCE_URL,
        source_attribution=NERO_ATTRIBUTION,
        source_version=source_version,
        file_path=file_path,
        parser=parse_nero_zip,
        content_type="application/zip",
        worker_id=worker_id,
        now=now,
    )


def run_ivi_import(
    repository: JobRepository,
    store: SnapshotStore,
    database_url: str,
    *,
    source_version: str,
    file_path: Path,
    worker_id: str = "jsa-ivi-importer",
    now: datetime | None = None,
) -> JsaImportResult:
    return _run_jsa_import(
        repository,
        store,
        database_url,
        dataset="ivi",
        source_key=IVI_SOURCE_KEY,
        source_name="Jobs and Skills Australia Internet Vacancy Index",
        source_url=IVI_SOURCE_URL,
        source_attribution=IVI_ATTRIBUTION,
        source_version=source_version,
        file_path=file_path,
        parser=parse_ivi_workbook,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        worker_id=worker_id,
        now=now,
    )
